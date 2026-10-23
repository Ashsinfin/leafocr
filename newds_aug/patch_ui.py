"""
Patch: adds TXT/PDF/JSON/CSV/DOCX export to app.py.
Keeps lines 1-1269 intact, rewrites the UI section (10.) only.
"""
import os

APP = os.path.join(os.path.dirname(__file__), "app.py")

with open(APP, "r", encoding="utf-8") as f:
    lines = f.readlines()

kept = lines[:1269]   # everything up to but NOT including line 1270 (the UI comment)

NEW_UI = r'''# ============================================================
# 10. GRADIO UI
# ============================================================
import csv as _csv
import json as _json
import io
import tempfile
import datetime

# ── Export helpers ───────────────────────────────────────────────────────────

def _timestamp():
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def export_txt(session):
    """Export plain text transcript."""
    if not session or not session.get("records"):
        return None
    text = plain_text(session["records"])
    path = os.path.join(tempfile.gettempdir(), f"leafocr_{_timestamp()}.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    return path


def export_json(session):
    """Export structured JSON with per-chunk metadata."""
    if not session or not session.get("records"):
        return None
    records = session["records"]
    data = {
        "exported_at": datetime.datetime.now().isoformat(),
        "total_lines": len({r["line"] for r in records}),
        "total_chunks": len(records),
        "chunks": [
            {
                "line": r["line"],
                "chunk": r["chunk"],
                "text": r["text"],
                "bbox": {"x1": r["rect"][0], "y1": r["rect"][1],
                         "x2": r["rect"][2], "y2": r["rect"][3]},
            }
            for r in records
        ],
        "full_text": plain_text(records),
    }
    path = os.path.join(tempfile.gettempdir(), f"leafocr_{_timestamp()}.json")
    with open(path, "w", encoding="utf-8") as f:
        _json.dump(data, f, ensure_ascii=False, indent=2)
    return path


def export_csv(session):
    """Export tabular CSV: line, chunk, text, x1, y1, x2, y2."""
    if not session or not session.get("records"):
        return None
    path = os.path.join(tempfile.gettempdir(), f"leafocr_{_timestamp()}.csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = _csv.writer(f)
        w.writerow(["line", "chunk", "text", "x1", "y1", "x2", "y2"])
        for r in session["records"]:
            x1, y1, x2, y2 = r["rect"]
            w.writerow([r["line"], r["chunk"], r["text"], x1, y1, x2, y2])
    return path


def export_docx(session):
    """Export a Word document with one paragraph per line."""
    if not session or not session.get("records"):
        return None
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    # Page margins
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(2.5)

    title = doc.add_heading("LeafOCR Transcript", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(0x25, 0x3e, 0x34)

    ts = doc.add_paragraph(
        f"Exported: {datetime.datetime.now().strftime('%d %B %Y, %H:%M')}"
    )
    ts.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ts.runs[0].font.size = Pt(9)
    ts.runs[0].font.color.rgb = RGBColor(0x7b, 0x86, 0x7b)
    doc.add_paragraph()   # spacer

    records = session["records"]
    lines_map = {}
    for r in records:
        lines_map.setdefault(r["line"], []).append(r["text"])

    for li in sorted(lines_map):
        para = doc.add_paragraph(" ".join(lines_map[li]))
        para.runs[0].font.size = Pt(14)
        para.paragraph_format.space_after = Pt(8)

    path = os.path.join(tempfile.gettempdir(), f"leafocr_{_timestamp()}.docx")
    doc.save(path)
    return path


def export_pdf(session):
    """Export a PDF using fpdf2 with the Nirmala TTC font (supports Malayalam)."""
    if not session or not session.get("records"):
        return None
    from fpdf import FPDF

    FONT_PATH = r"C:\Windows\Fonts\Nirmala.ttc"
    FALLBACK  = None
    # Check if Malayalam font is available
    has_font = os.path.exists(FONT_PATH)

    class PDF(FPDF):
        pass

    pdf = PDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_margins(25, 20, 20)

    if has_font:
        pdf.add_font("Nirmala", style="", fname=FONT_PATH)
        pdf.add_font("Nirmala", style="B", fname=FONT_PATH)
        body_font = "Nirmala"
    else:
        body_font = "Helvetica"

    # Title
    pdf.set_font(body_font, style="B", size=18)
    pdf.set_text_color(37, 62, 52)
    pdf.cell(0, 10, "LeafOCR Transcript", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    # Timestamp
    pdf.set_font(body_font, size=9)
    pdf.set_text_color(123, 134, 123)
    pdf.cell(0, 6,
             f"Exported: {datetime.datetime.now().strftime('%d %B %Y, %H:%M')}",
             align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    # Divider
    pdf.set_draw_color(217, 223, 212)
    pdf.line(25, pdf.get_y(), 190, pdf.get_y())
    pdf.ln(8)

    # Body text
    pdf.set_font(body_font, size=14)
    pdf.set_text_color(52, 69, 54)

    records = session["records"]
    lines_map = {}
    for r in records:
        lines_map.setdefault(r["line"], []).append(r["text"])

    for li in sorted(lines_map):
        line_text = " ".join(lines_map[li])
        pdf.multi_cell(0, 9, line_text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

    # Footer
    pdf.set_y(-18)
    pdf.set_font(body_font, size=8)
    pdf.set_text_color(160, 170, 155)
    pdf.cell(0, 6, "Generated by LeafOCR · Palm-leaf manuscript recognition", align="C")

    path = os.path.join(tempfile.gettempdir(), f"leafocr_{_timestamp()}.pdf")
    pdf.output(path)
    return path


# Dispatcher — called by the export button
_EXPORT_FNS = {
    "Plain text (.txt)":   export_txt,
    "JSON (.json)":        export_json,
    "CSV (.csv)":          export_csv,
    "Word document (.docx)": export_docx,
    "PDF (.pdf)":          export_pdf,
}

def do_export(fmt, session):
    fn = _EXPORT_FNS.get(fmt)
    if fn is None or not session or not session.get("records"):
        return None
    try:
        return fn(session)
    except Exception as e:
        print(f"[Export error] {fmt}: {e}")
        return None


# ── CSS ──────────────────────────────────────────────────────────────────────

APP_CSS = """
body, .dark, gradio-app { background: #f3f2ec !important; }
.gradio-container {
    max-width: 1500px !important; margin: 0 auto !important; padding: 24px 32px !important;
    font-family: 'Segoe UI', sans-serif !important; color: #253e34 !important;
    --body-background-fill: #f3f2ec; --body-text-color: #253e34;
    --body-text-color-subdued: #69766d; --background-fill-primary: #fffefa;
    --background-fill-secondary: #f1f3ed; --block-background-fill: #fffefa;
    --block-border-color: #dce2d8; --border-color-primary: #dce2d8;
    --border-color-accent: #537963; --input-background-fill: #f7f8f3;
    --input-border-color: #dce2d8; --input-placeholder-color: #7b877e;
    --block-label-background-fill: #fffefa; --block-label-text-color: #52675a;
    --block-title-text-color: #253e34; --input-text-color: #253e34;
    --color-accent: #346b50; --checkbox-background-color-selected: #346b50;
    --button-secondary-background-fill: #eff3ed; --button-secondary-text-color: #253e34;
}
.gradio-container [data-testid="block-info"] { background:transparent !important; color:#65765f !important; font-size:11px !important; font-weight:600 !important; box-shadow:none !important; }
.brandbar { display:flex; align-items:center; justify-content:space-between; padding:0 0 20px; border-bottom:1px solid #d9dfd4; }
.brand { display:flex; align-items:center; gap:10px; font-weight:700; font-size:23px; letter-spacing:-1px; }
.brand-mark { display:grid; place-items:center; width:38px; height:38px; background:#214d38; color:#e7efd9; border-radius:12px 12px 12px 3px; }
.brand-note { color:#728071; font-size:11px; text-transform:uppercase; letter-spacing:2px; }
.brand-right { display:flex; align-items:center; gap:20px; }
.leaf-seal { width:76px; height:76px; flex:none; border:1px solid #d1dac9; border-radius:50%; display:grid; place-content:center; text-align:center; color:#708368; font:italic 30px Georgia,serif; background:radial-gradient(circle,#e8edde,#f3f2ec); }
.leaf-seal small { font:7px 'Segoe UI',sans-serif; letter-spacing:1.1px; margin-top:4px; }

/* ── Scanner toggle toolbar ── */
#scanner-toolbar { display:flex; align-items:center; justify-content:flex-end; gap:10px; margin-bottom:10px; }
#scanner-toolbar > div { flex:none !important; width:auto !important; }
#toggle-scanner {
    display:inline-flex !important; align-items:center !important; gap:6px !important;
    background:#eff3ed !important; border:1px solid #cfd8c7 !important; border-radius:9px !important;
    padding:6px 14px !important; font-size:12px !important; font-weight:600 !important;
    color:#3d5c44 !important; cursor:pointer !important;
    transition:background .15s, box-shadow .15s !important;
    white-space:nowrap !important; min-height:0 !important; height:34px !important;
    width:auto !important;
}
#toggle-scanner:hover { background:#e3ead9 !important; box-shadow:0 2px 6px #2a4a3510 !important; }

/* ── Stage cards ── */
.stage-card {
    background:#fffefa !important; border:1px solid #dce2d8 !important;
    border-radius:18px !important; padding:20px !important;
    box-shadow:0 5px 20px #233e2f06; gap:14px !important;
}
.stage-title { height:78px; flex:none !important; }
.stage-heading { display:flex; align-items:center; gap:10px; margin-bottom:9px; }
.step { display:grid; place-items:center; width:28px; height:28px; border-radius:9px; background:#edf2e7; color:#57734d; font-size:11px; font-weight:700; }
.stage-heading h2 { font-size:17px; font-weight:650; color:#253e34; margin:0; letter-spacing:-.3px; }
.stage-title p { color:#7b867b; font-size:12px; line-height:1.65; margin:0; }
.status-bar { background:#e9efe2; border:1px solid #d7e1ce; border-radius:10px; padding:9px 14px !important; margin-bottom:6px; }
.status-bar p { font-size:12px !important; color:#50674a !important; margin:0 !important; }
.status-bar p:before { content:'●'; color:#598152; margin-right:9px; font-size:10px; }

/* ── Workspace row ── */
#workspace { gap:20px; align-items:stretch !important; }

/* ── Rectangular image panel ── */
.image-stage {
    height:360px !important; min-height:320px !important; flex:none !important;
    border:1px solid #cfd8c7 !important; border-radius:12px !important;
    background:#f3f5ee !important; overflow:hidden !important;
    display:flex !important; align-items:center !important; justify-content:center !important;
}
#ms-upload .image-stage { border-style:dashed !important; }
.image-stage img { object-fit:contain !important; object-position:center !important; width:100% !important; height:100% !important; }
.image-stage .image-container { background:transparent !important; width:100% !important; height:100% !important; }

/* ── Rectangular transcript panel ── */
#transcript {
    height:300px !important; overflow-y:auto !important; flex:none !important;
    background:#faf8f0 !important; border:1px solid #e8e3d4 !important;
    border-radius:12px !important; box-sizing:border-box !important;
}
#transcript, #transcript * { font-family:'Nirmala UI','Noto Sans Malayalam',sans-serif; }
#transcript .textfield { font-size:19px; line-height:1.85; color:#344536; }
#transcript .token-container { white-space:nowrap; }
#transcript .token-container:empty { flex-basis:100%; height:0; }
#transcript .token { margin:2px 0 !important; padding:0 4px !important; cursor:pointer; border-radius:5px; }
#transcript .token:not(.highlighted) { padding:0 !important; cursor:default; }
#transcript .token.highlighted:hover { filter:brightness(.92); outline:1px solid #b9c7ad; }
#transcript > label { background:transparent !important; box-shadow:none !important; border:0 !important; }

/* ── Source region panel ── */
#region { height:110px !important; flex:none !important; border-radius:12px !important; }

/* ── Controls ── */
#run-controls { gap:8px; }
#copy-text { align-self:flex-end; flex:none !important; min-height:0 !important; height:30px; width:auto !important; padding:0 12px !important; border-radius:9px !important; font-size:12px !important; }
#run-controls button { border-radius:9px !important; font-size:12px !important; }
#recognize { background:#28563e !important; color:#fffef8 !important; border:1px solid #28563e !important; border-radius:10px !important; font-size:14px !important; font-weight:600 !important; min-height:46px; box-shadow:0 3px 8px #28563e20; transition:background .15s,transform .15s; }
#recognize:hover { background:#183e2a !important; transform:translateY(-1px); }
.panel-note { color:#81907e; font-size:11px; line-height:1.7; padding-top:3px; }
#details { margin-top:10px; background:#fffefa !important; border:1px solid #dce2d8 !important; border-radius:12px !important; }
#details button, #details button span { color:#53654d !important; }
.app-foot { display:flex; justify-content:space-between; gap:16px; color:#8a9485; font-size:10px; letter-spacing:.7px; padding:14px 2px 4px; }
footer { display:none !important; }

/* ── Export bar ── */
#export-bar {
    display:flex; align-items:center; gap:8px;
    background:#f3f5ee; border:1px solid #dce2d8; border-radius:11px;
    padding:8px 12px; margin-top:4px;
}
#export-bar > * { margin:0 !important; }
#export-bar label { display:none !important; }        /* hide Gradio label */
#export-fmt { flex:1 1 auto; min-width:0 !important; }
#export-fmt select { font-size:12px !important; height:32px !important; border-radius:8px !important; padding:0 8px !important; }
#export-btn {
    flex:none !important; min-width:0 !important; width:auto !important;
    height:32px !important; min-height:0 !important;
    padding:0 14px !important; font-size:12px !important;
    border-radius:8px !important; font-weight:600 !important;
    background:#28563e !important; color:#fffef8 !important;
    border:1px solid #28563e !important;
    transition:background .15s !important;
}
#export-btn:hover { background:#183e2a !important; }
/* Download link that Gradio renders after export */
#export-file { flex:none; }
#export-file a { font-size:11px !important; color:#346b50 !important; font-weight:600; text-decoration:none; }
#export-file a:hover { text-decoration:underline; }

/* ── Loading skeleton ── */
#ocr-loading:not(:has(.ocr-skeleton)) { display:none !important; }
#workspace:has(.ocr-skeleton) #transcript { display:none !important; }
.ocr-skeleton { height:300px; box-sizing:border-box; padding:22px 18px; display:flex; flex-direction:column; gap:18px; background:#faf8f0; border:1px solid #e8e3d4; border-radius:12px; overflow:hidden; animation:sk-enter .3s ease-out both; }
.sk-lines { display:flex; flex-direction:column; gap:11px; }
.sk-line { position:relative; display:block; height:10px; flex:none; border-radius:5px; background:#e8e6d9; overflow:hidden; }
.sk-line::after { content:''; position:absolute; inset:0; background:linear-gradient(100deg,transparent 15%,#fffdf4b3 48%,transparent 80%); transform:translateX(-110%); animation:sk-shimmer 2.8s cubic-bezier(.4,0,.2,1) infinite; animation-delay:var(--sk-delay,0s); }
.sk-caption { margin-top:auto; display:flex; align-items:center; gap:9px; color:#71816b; font-size:11px; letter-spacing:.25px; line-height:18px; }
.sk-dots { display:inline-flex; gap:4px; align-items:center; }
.sk-dots i { width:4px; height:4px; border-radius:50%; background:#779369; animation:sk-breathe 1.8s ease-in-out infinite; }
.sk-dots i:nth-child(2) { animation-delay:.18s; }
.sk-dots i:nth-child(3) { animation-delay:.36s; }
.live-scan { position:relative; overflow:hidden; }
#workspace:has(.ocr-skeleton) .live-scan::after { content:''; position:absolute; inset:0; pointer-events:none; z-index:5;
    background:linear-gradient(180deg,transparent 28%,#7fc98a08 40%,#7fc98a26 49.5%,#59815299 50%,transparent 50.5%); animation:sk-sweep 3.6s cubic-bezier(.4,0,.2,1) infinite; }
@keyframes sk-enter { from { opacity:0; } to { opacity:1; } }
@keyframes sk-shimmer { 0%,12% { transform:translateX(-110%); } 72%,100% { transform:translateX(110%); } }
@keyframes sk-breathe { 0%,70%,100% { opacity:.4; transform:translateY(0); } 35% { opacity:1; transform:translateY(-2px); } }
@keyframes sk-sweep { 0% { transform:translateY(-65%); opacity:0; } 12%,80% { opacity:1; } 100% { transform:translateY(65%); opacity:0; } }
@media(prefers-reduced-motion:reduce) {
    .ocr-skeleton, .sk-line::after, .sk-dots i, #workspace:has(.ocr-skeleton) .live-scan::after { animation:none !important; }
    .sk-line::after, #workspace:has(.ocr-skeleton) .live-scan::after { display:none !important; }
}

/* ── Responsive ── */
.gradio-container.gradio-container { width:100% !important; min-width:0 !important; box-sizing:border-box; padding:clamp(12px,3vw,32px) !important; }
.gradio-container .main { width:100% !important; min-width:0 !important; padding:0 !important; box-sizing:border-box; }

/* Tablet 601-960px */
@media(min-width:601px) and (max-width:960px) {
    #workspace { flex-wrap:wrap; gap:16px; }
    #workspace > div:nth-child(1),
    #workspace > div:nth-child(3) { flex:1 1 calc(50% - 10px) !important; min-width:240px !important; box-sizing:border-box; }
    #scanner-col { flex:1 1 100% !important; min-width:0 !important; order:99; }
    #scanner-col .image-stage { height:240px !important; min-height:200px !important; }
    .image-stage { height:300px !important; min-height:260px !important; }
    #transcript { height:250px !important; }
    #workspace .stage-title { height:auto; min-height:64px; }
    #export-bar { flex-wrap:wrap; }
}

/* Mobile <=600px */
@media(max-width:600px) {
    #workspace { flex-direction:column !important; gap:14px; }
    #workspace .stage-card { min-width:0 !important; width:100% !important; box-sizing:border-box; padding:16px !important; }
    #scanner-col { order:99; }
    .brandbar .brand-note { display:none; }
    .brandbar .leaf-seal { width:60px; height:60px; font-size:25px; }
    .brandbar { gap:12px; }
    #workspace .stage-title { height:auto; min-height:64px; }
    #workspace .image-stage { height:260px !important; min-height:220px !important; }
    #transcript { height:220px !important; }
    .app-foot { flex-wrap:wrap; gap:6px; }
    #scanner-toolbar { justify-content:flex-start; }
    #export-bar { flex-wrap:wrap; }
    #export-fmt { min-width:100% !important; }
}
"""


def _toggle_scanner(is_visible):
    """Toggle the Live Scanner column visibility."""
    new_visible = not is_visible
    label = "\u25be Hide Live Scanner" if new_visible else "\u25b4 Show Live Scanner"
    return new_visible, gr.update(visible=new_visible), gr.update(value=label)


with gr.Blocks(title="Palm Leaf OCR") as demo:
    session_state   = gr.State(new_session())
    scanner_visible = gr.State(False)   # scanner starts collapsed

    gr.HTML(
        '<div class="brandbar"><div class="brand"><span class="brand-mark">\u2381</span>LeafOCR</div>'
        '<div class="brand-right"><div class="brand-note">Palm-leaf manuscript studio</div>'
        '<div class="leaf-seal" aria-label="Leaf to letter emblem">\u0d05<small>LEAF TO LETTER</small></div></div></div>'
    )

    status_output = gr.Markdown("Ready \u00b7 upload an image to begin", elem_classes="status-bar")

    # ── Scanner toggle toolbar ──────────────────────────────────────────────
    with gr.Row(elem_id="scanner-toolbar"):
        toggle_btn = gr.Button(
            "\u25b4 Show Live Scanner",
            size="sm",
            variant="secondary",
            elem_id="toggle-scanner",
            scale=0,
        )

    # ── Three-column workspace ─────────────────────────────────────────────
    with gr.Row(equal_height=True, elem_id="workspace"):

        # Column 1 — Your manuscript
        with gr.Column(scale=2, min_width=280, elem_classes="stage-card", elem_id="ms-upload"):
            gr.HTML(
                '<div class="stage-heading"><span class="step">01</span>'
                '<h2>Your manuscript</h2></div>'
                '<p>Start with a leaf image or a cropped text segment.</p>',
                elem_classes="stage-title",
            )
            input_image = gr.Image(
                label="Source image", type="numpy", height=360,
                scale=0, elem_classes="image-stage",
                buttons=["fullscreen"], sources=["upload", "clipboard"],
            )
            run_btn = gr.Button(
                "Recognize manuscript  \u2192",
                variant="primary", size="lg", elem_id="recognize",
            )
            with gr.Row(elem_id="run-controls"):
                stop_btn   = gr.Button("\u25a0 Stop",   size="sm", min_width=60)
                resume_btn = gr.Button("\u25b6 Resume", size="sm", min_width=60, visible=False)
                reset_btn  = gr.Button("\u21ba Reset",  size="sm", min_width=60)

        # Column 2 — Live scanner (starts hidden)
        with gr.Column(
            scale=2, min_width=260,
            elem_classes="stage-card", elem_id="scanner-col",
            visible=False,
        ) as scanner_col:
            gr.HTML(
                '<div class="stage-heading"><span class="step">02</span>'
                '<h2>Live scanner</h2></div>'
                '<p>Follow each text region as it is recognized.</p>',
                elem_classes="stage-title",
            )
            annotated_output = gr.Image(
                label="Live scan", height=360, interactive=False,
                scale=0, elem_classes=["image-stage", "live-scan"],
                buttons=["fullscreen", "download"],
            )
            gr.HTML(
                '<div class="panel-note">'
                '<b style="color:#2f8a3c">Green</b> marks the text being read right now; '
                '<b style="color:#d9822b">orange</b> marks text you picked in the transcript.<br>'
                'Open full screen to inspect the finer details.</div>'
            )

        # Column 3 — The transcript
        with gr.Column(scale=2, min_width=280, elem_classes="stage-card"):
            gr.HTML(
                '<div class="stage-heading"><span class="step">03</span>'
                '<h2>The transcript</h2></div>'
                '<p>Click any word to see where it came from.</p>',
                elem_classes="stage-title",
            )
            skeleton_output = gr.HTML("", elem_id="ocr-loading")
            output_text = gr.HighlightedText(
                label="Recognized text",
                color_map={"chunk": "#f1efe3", "current": "#7fc98a", "selected": "#f5a54a"},
                show_legend=False,
                show_inline_category=False,
                interactive=False,
                elem_id="transcript",
            )

            # ── Transcript action row: Copy + Export ──────────────────────
            with gr.Row(elem_classes="transcript-actions"):
                copy_btn = gr.Button(
                    "\u29c9 Copy text", size="sm", variant="secondary",
                    scale=0, elem_id="copy-text",
                )

            # ── Export bar ────────────────────────────────────────────────
            with gr.Row(elem_id="export-bar"):
                export_fmt = gr.Dropdown(
                    choices=list(_EXPORT_FNS.keys()),
                    value="Plain text (.txt)",
                    label="Format",
                    container=False,
                    scale=3,
                    elem_id="export-fmt",
                )
                export_btn = gr.Button(
                    "\u2913 Export",
                    size="sm",
                    scale=1,
                    elem_id="export-btn",
                )
                export_file = gr.File(
                    label="",
                    visible=False,
                    elem_id="export-file",
                    scale=2,
                )

            region_output = gr.Image(
                label="Source region", height=110, interactive=False,
                elem_id="region", buttons=["fullscreen"],
            )

    with gr.Accordion(
        "Behind the scan \u00b7 lines & text regions",
        open=False, elem_id="details",
    ):
        plain_output  = gr.Textbox(label="Plain text", lines=4, buttons=["copy"], interactive=False)
        mask_output   = gr.Image(label="Segmentation analysis")
        line_output   = gr.Gallery(label="Detected lines", columns=1, object_fit="contain")
        chunk_output  = gr.Gallery(
            label="Recognized chunks (click one to locate it)",
            columns=4, object_fit="contain",
        )

    gr.HTML(
        '<div class="app-foot">'
        '<span>LeafOCR \u00b7 Manuscript recognition</span>'
        '<span>From fragile leaves to lasting words.</span>'
        '</div>'
    )

    # ── Output list (must match OUT_KEYS order) ────────────────────────────
    all_outputs = [
        mask_output, annotated_output, line_output, chunk_output, output_text,
        plain_output, status_output, region_output, session_state, stop_btn, resume_btn,
        skeleton_output,
    ]

    # ── Pipeline events (unchanged) ────────────────────────────────────────
    run_evt    = run_btn.click(fn=run_pipeline,    inputs=[input_image, session_state], outputs=all_outputs)
    resume_evt = resume_btn.click(fn=resume_pipeline, inputs=[input_image, session_state], outputs=all_outputs)
    running    = [run_evt, resume_evt]

    stop_btn.click(fn=stop_pipeline, inputs=[session_state], outputs=all_outputs,
                   cancels=running, concurrency_limit=None)
    reset_btn.click(fn=reset_session, inputs=[session_state], outputs=all_outputs,
                    cancels=running, concurrency_limit=None)
    input_image.upload(fn=reset_session, inputs=[session_state], outputs=all_outputs,
                       cancels=running, concurrency_limit=None)
    input_image.clear(fn=reset_session, inputs=[session_state], outputs=all_outputs,
                      cancels=running, concurrency_limit=None)

    output_text.select(fn=select_from_transcript, inputs=[session_state], outputs=all_outputs,
                       concurrency_limit=None)
    chunk_output.select(fn=select_from_gallery, inputs=[session_state], outputs=all_outputs,
                        concurrency_limit=None)
    copy_btn.click(fn=None, inputs=[plain_output],
                   js="(t) => { navigator.clipboard.writeText(t || ''); }")

    # ── Export event ───────────────────────────────────────────────────────
    def _do_export_ui(fmt, session):
        path = do_export(fmt, session)
        if path:
            return gr.update(value=path, visible=True)
        return gr.update(value=None, visible=False)

    export_btn.click(
        fn=_do_export_ui,
        inputs=[export_fmt, session_state],
        outputs=[export_file],
        concurrency_limit=None,
    )

    # ── Scanner toggle ─────────────────────────────────────────────────────
    toggle_btn.click(
        fn=_toggle_scanner,
        inputs=[scanner_visible],
        outputs=[scanner_visible, scanner_col, toggle_btn],
    )


if __name__ == "__main__":
    demo.queue(default_concurrency_limit=1).launch(
        theme=gr.themes.Soft(primary_hue="green", neutral_hue="stone"),
        css=APP_CSS,
    )
'''

with open(APP, "w", encoding="utf-8") as f:
    f.writelines(kept)
    f.write(NEW_UI)

print(f"Patched {APP}  ({len(kept)} lines kept + {len(NEW_UI)} chars new UI)")
